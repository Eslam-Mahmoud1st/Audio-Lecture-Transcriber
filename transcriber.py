import os
import whisper
import numpy as np
from docx import Document
from scipy.io import wavfile
from moviepy.audio.io.AudioFileClip import AudioFileClip
from moviepy.video.io.VideoFileClip import VideoFileClip

def extract_and_load_audio(input_path):

    temp_audio = "temp_lecture_audio.wav"
    print("[*] Extracting audio track...")
    
    try:
        clip = AudioFileClip(input_path)
        clip.write_audiofile(temp_audio, fps=16000, nbytes=2, codec='pcm_s16le')
        clip.close()
    except Exception:
        try:
            video = VideoFileClip(input_path)
            if video.audio is not None:
                video.audio.write_audiofile(temp_audio, fps=16000, nbytes=2, codec='pcm_s16le')
                video.close()
            else:
                video.close()
                return None
        except Exception as e:
            print(f"[!] Could not extract audio: {e}")
            return None

    try:
        sample_rate, data = wavfile.read(temp_audio)
        
        if len(data.shape) > 1:
            data = data.mean(axis=1)
            
        audio_array = data.astype(np.float32) / 32768.0

        if os.path.exists(temp_audio):
            os.remove(temp_audio)

        return audio_array
    except Exception as e:
        print(f"[!] Failed to read WAV array: {e}")
        if os.path.exists(temp_audio):
            os.remove(temp_audio)
        return None

def transcribe_media(file_path, model_size="base"):
    if not os.path.exists(file_path):
        print(f"[!] File not found: {file_path}")
        return None

    audio_array = extract_and_load_audio(file_path)
    if audio_array is None:
        print("[-] Audio extraction failed.")
        return None

    print(f"[*] Loading Whisper model ('{model_size}')...")
    model = whisper.load_model(model_size)

    print(f"[*] Transcribing text... This may take a few moments.")
    result = model.transcribe(audio_array, fp16=False)

    return result.get("text", "")

def save_to_docx(text, output_filename):
    doc = Document()
    doc.add_heading("Lecture / Video Notes & Transcription", level=1)
    doc.add_paragraph(text)
    doc.save(output_filename)
    print(f"[✓] Document saved successfully to '{output_filename}'")

def main():
    print("=== Audio & Video Lecture Transcriber ===")
    media_file = input("Enter path to audio/video file (e.g. lecture.mp4, audio.mp3): ").strip()
    media_file = media_file.strip('"').strip("'")

    if not media_file:
        print("[!] No file path provided.")
        return

    transcribed_text = transcribe_media(media_file, model_size="base")

    if transcribed_text:
        print("\n--- Transcription Preview ---")
        print(transcribed_text[:300] + ("..." if len(transcribed_text) > 300 else ""))
        print("-----------------------------\n")

        base_name = os.path.splitext(os.path.basename(media_file))[0]
        output_docx = f"{base_name}_notes.docx"
        save_to_docx(transcribed_text, output_docx)
    else:
        print("[-] Processing failed.")

if __name__ == "__main__":
    main()